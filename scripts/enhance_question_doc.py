from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

src = '/Users/linglingqi/Desktop/秋招群面题目分类.docx'
out = '/Users/linglingqi/Desktop/秋招群面题目分类_完善版.docx'
doc = Document(src)

for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    if style_name in doc.styles:
        doc.styles[style_name].font.name = 'Hiragino Sans GB'
        doc.styles[style_name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
for para in doc.paragraphs:
    for run in para.runs:
        run.font.name = 'Hiragino Sans GB'
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
doc.add_page_break()

def h(text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(12)
    p.add_run(text)

def p(text):
    x = doc.add_paragraph(text)
    x.paragraph_format.space_after = Pt(6)
    x.paragraph_format.line_spacing = 1.15

h('五 首期产品可执行分类标准')
p('用户进入练习时按行业、岗位、题型依次选择；系统从匹配题库中取题，并根据题型从 12 张角色卡中动态分配 3 名 AI 候选人。')
h('5.1 行业范围', 2)
p('互联网：产品经理、产品运营、用户增长、商业分析、管培生。重点考察用户视角、逻辑拆解、数据分析、产品落地和资源取舍。')
p('快消：品牌市场、销售渠道、产品运营、商业分析、管培生。重点考察消费者洞察、品牌策略、渠道管理、ROI 和商业落地。')
p('金融：银行、券商、基金资管、金融产品、商业分析、管培生。重点考察风险合规、收益权衡、客户分析、资产配置和宏观判断。')
p('央国企：综合管培、行政党群、人力、工程产业、投融资。重点考察大局观、执行力、合规意识、社会责任和稳定性。')
h('5.2 岗位范围', 2)
p('首期岗位建议统一为：产品经理、产品运营 / 用户增长、品牌市场 / 销售渠道、商业分析、管理培训生 / 综合管培、工程产业 / 投融资方向。岗位可绑定多个行业，题目通过适用岗位字段进行筛选。')
h('5.3 首期题型', 2)
p('排序题：在资源、时间或人力有限时，对多个事项、需求或项目排序，并说明判断标准。评价重点是标准清晰、取舍有据、优先级稳定和推进共识。')
p('方案设计题：围绕目标设计可落地方案，说明目标人群、策略路径、资源投入、指标和风险。评价重点是结构完整、执行可行和指标闭环。')
p('开放辩论题：围绕社会议题、商业取舍或价值冲突表达立场，比较不同方案并形成结论。评价重点是观点独立、论据充分、回应他人和兼顾多方。')
h('六 题目来源与使用规则')
p('题目来源优先级保持不变：第一优先使用静态题库中的已审核题目；第二优先使用题型模板结合行业、岗位和目标公司生成题目；只有前两级不可用时，才调用 AI 生成兜底题目。进入面试前保存最终题目文本、行业、岗位、题型和来源标记。')
p('静态题库至少记录：题目编号、行业、适用岗位、题型、难度、题目正文、约束条件、考察要点、参考讨论方向和版本状态。')
h('七 题型与八阶段流程', 1)
p('自我介绍：3 名 AI 候选人与用户依次介绍。读题和独立思考：展示题目、倒计时和思考提示，不触发自由讨论。小组讨论：3 名 AI 候选人与用户并行进行语音和文字交流。')
p('汇报：形成团队结论并完成汇报；排序题输出排序结果，方案题输出方案闭环，辩论题输出共识与保留分歧。面试官提问：优先追问排序依据、方案取舍、观点边界和用户个人贡献。反问和复盘：用户向面试官提问，系统保存完整过程并生成报告。')
h('八 首期题库建设建议', 1)
p('首期建议先完成 12 道可直接开练的标准题：每个行业覆盖 1 道排序题、1 道方案设计题和 1 道开放辩论题。每道题绑定 1 至 3 个岗位，确保完成“行业 → 岗位 → 题型”选择后能够直接进入面试。')
p('题库上线前逐题检查：任务是否明确；约束是否足够支持讨论；题型标签是否准确；是否能在八阶段流程中形成可评价的团队结论。')
h('九 推荐的产品筛选路径', 1)
p('匿名用户进入准备页后，依次选择行业、岗位和题型。系统优先展示题库题目，确认后进入设备检测和群面房间。房间固定展示 3 名 AI 候选人、1 名 AI 面试官和用户；结束后生成包含聊天、转写、阶段时间、发言事件和报告结果的完整练习记录。报告使用本场表现位次、模拟排名和表现等级，不表述为现实招聘结果预测。')

doc.save(out)
print(out)
